#!/usr/bin/env python3
"""
RCC 추출 데이터의 KC 인증정보 유효성 검증

입력: CSV/Excel (details 칼럼에 인증번호·제품명·모델명 포함)
검증: SafetyKorea searchPop API
출력: VALID / INVALID / MISMATCH / NOT_FOUND CSV 리포트

사용법:
  python scripts/validate_kc_certification.py \\
    --input samples/kc_validation_input.csv \\
    --output reports/kc_validation_report.csv \\
    --insecure

  python scripts/validate_kc_certification.py \\
    --input data/rcc_items.xlsx \\
    --output reports/kc_validation_report.csv \\
    --details-column details
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from kc_safetykorea_client import KCCertDetail, fetch_cert_detail

CERT_NUM_KEYS = (
    "cert_num",
    "certNum",
    "certification_number",
    "certificationNumber",
    "auth_num",
    "authNum",
    "kc_cert_num",
    "인증번호",
)
PRODUCT_KEYS = ("product_name", "productName", "product", "item_name", "제품명")
MODEL_KEYS = ("model_name", "modelName", "model", "모델명")
KC_CERT_PATTERN = re.compile(r"\b([A-Z]{2}\d{3,6}-[\dA-Z]+)\b", re.IGNORECASE)

VALIDATION_VALID = "VALID"
VALIDATION_INVALID = "INVALID"
VALIDATION_MISMATCH = "MISMATCH"
VALIDATION_NOT_FOUND = "NOT_FOUND"
VALIDATION_SKIPPED = "SKIPPED"


@dataclass
class InputCertRecord:
    row_number: int
    cert_num: str
    product_name: str
    model_name: str
    details_raw: str
    source_row: dict[str, str] = field(default_factory=dict)


@dataclass
class ValidationResult:
    row_number: int
    validation_status: str
    validation_message: str
    cert_num_input: str
    product_name_input: str
    model_name_input: str
    cert_num_api: str
    product_name_api: str
    model_name_api: str
    cert_status_api: str
    recall_status_api: str
    cert_date_api: str
    cert_org_api: str
    source_row: dict[str, str] = field(default_factory=dict)


def normalize_text(value: str) -> str:
    text = (value or "").strip().lower()
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[·•\-_/]", "", text)
    return text


def normalize_product_name(value: str) -> str:
    text = (value or "").strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def names_match(input_value: str, api_value: str, *, strict: bool = True) -> bool:
    if not input_value:
        return True
    if not api_value:
        return False

    if strict:
        return normalize_text(input_value) == normalize_text(api_value)

    input_norm = normalize_product_name(input_value)
    api_norm = normalize_product_name(api_value)
    if input_norm == api_norm:
        return True
    return input_norm in api_norm or api_norm in input_norm


def has_recall(recall_status: str) -> bool:
    value = (recall_status or "").strip()
    if not value:
        return False
    return value not in {"-", "없음", "해당없음", "N", "NO", "none"}


def _pick_from_mapping(data: dict[str, Any], keys: tuple[str, ...]) -> str:
    for key in keys:
        if key in data and data[key] not in (None, ""):
            return str(data[key]).strip()
    return ""


def _extract_from_object(obj: Any) -> dict[str, str]:
    if not isinstance(obj, dict):
        return {}

    cert_num = _pick_from_mapping(obj, CERT_NUM_KEYS)
    product_name = _pick_from_mapping(obj, PRODUCT_KEYS)
    model_name = _pick_from_mapping(obj, MODEL_KEYS)

    for value in obj.values():
        if isinstance(value, dict):
            nested = _extract_from_object(value)
            cert_num = cert_num or nested.get("cert_num", "")
            product_name = product_name or nested.get("product_name", "")
            model_name = model_name or nested.get("model_name", "")

    return {
        "cert_num": cert_num,
        "product_name": product_name,
        "model_name": model_name,
    }


def parse_details(details: str) -> dict[str, str]:
    text = (details or "").strip()
    if not text:
        return {"cert_num": "", "product_name": "", "model_name": ""}

    if text.startswith("{") or text.startswith("["):
        try:
            payload = json.loads(text)
            if isinstance(payload, list):
                merged = {"cert_num": "", "product_name": "", "model_name": ""}
                for item in payload:
                    extracted = _extract_from_object(item)
                    for key in merged:
                        merged[key] = merged[key] or extracted.get(key, "")
                return merged
            return _extract_from_object(payload)
        except json.JSONDecodeError:
            pass

    cert_match = KC_CERT_PATTERN.search(text)
    cert_num = cert_match.group(1).upper() if cert_match else ""

    product_name = ""
    model_name = ""
    for pattern, target in (
        (r"(?:제품명|product_name|productName)\s*[:=]\s*([^,;\n|]+)", "product_name"),
        (r"(?:모델명|model_name|modelName)\s*[:=]\s*([^,;\n|]+)", "model_name"),
    ):
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            if target == "product_name":
                product_name = match.group(1).strip()
            else:
                model_name = match.group(1).strip()

    return {
        "cert_num": cert_num,
        "product_name": product_name,
        "model_name": model_name,
    }


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def read_excel_rows(path: Path, sheet_name: Optional[str] = None) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError(
            "Excel 입력을 읽으려면 openpyxl이 필요합니다: pip install openpyxl"
        ) from e

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    records: list[dict[str, str]] = []
    for row in rows:
        record = {
            header[i]: "" if i >= len(row) or row[i] is None else str(row[i]).strip()
            for i in range(len(header))
        }
        records.append(record)
    return records


def load_input_records(
    path: Path,
    *,
    details_column: str = "details",
) -> list[InputCertRecord]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = read_csv_rows(path)
    elif suffix in {".xlsx", ".xlsm", ".xls"}:
        rows = read_excel_rows(path)
    else:
        raise RuntimeError(f"지원하지 않는 입력 형식입니다: {suffix}")

    records: list[InputCertRecord] = []
    for index, row in enumerate(rows, start=2):
        details_raw = row.get(details_column, "")
        parsed = parse_details(details_raw)
        records.append(
            InputCertRecord(
                row_number=index,
                cert_num=parsed["cert_num"],
                product_name=parsed["product_name"],
                model_name=parsed["model_name"],
                details_raw=details_raw,
                source_row=row,
            )
        )
    return records


def validate_record(
    record: InputCertRecord,
    detail: KCCertDetail,
) -> ValidationResult:
    if not record.cert_num:
        return ValidationResult(
            row_number=record.row_number,
            validation_status=VALIDATION_SKIPPED,
            validation_message="details에서 인증번호를 추출하지 못했습니다.",
            cert_num_input="",
            product_name_input=record.product_name,
            model_name_input=record.model_name,
            cert_num_api="",
            product_name_api="",
            model_name_api="",
            cert_status_api="",
            recall_status_api="",
            cert_date_api="",
            cert_org_api="",
            source_row=record.source_row,
        )

    if not detail.found:
        return ValidationResult(
            row_number=record.row_number,
            validation_status=VALIDATION_NOT_FOUND,
            validation_message="SafetyKorea에 인증정보가 존재하지 않습니다.",
            cert_num_input=record.cert_num,
            product_name_input=record.product_name,
            model_name_input=record.model_name,
            cert_num_api="",
            product_name_api="",
            model_name_api="",
            cert_status_api="",
            recall_status_api="",
            cert_date_api="",
            cert_org_api="",
            source_row=record.source_row,
        )

    issues: list[str] = []
    mismatch_fields: list[str] = []

    if detail.cert_status != "적합":
        issues.append(f"인증상태가 적합이 아님 (API={detail.cert_status or '없음'})")

    if has_recall(detail.recall_status):
        issues.append(f"리콜 발생 (API={detail.recall_status})")

    if record.product_name and not names_match(record.product_name, detail.product_name, strict=False):
        mismatch_fields.append(
            f"product_name: input={record.product_name}, api={detail.product_name}"
        )

    if record.model_name and not names_match(record.model_name, detail.model_name, strict=True):
        mismatch_fields.append(
            f"model_name: input={record.model_name}, api={detail.model_name}"
        )

    if mismatch_fields:
        status = VALIDATION_MISMATCH
        message = "; ".join(mismatch_fields)
    elif issues:
        status = VALIDATION_INVALID
        message = "; ".join(issues)
    else:
        status = VALIDATION_VALID
        message = "인증번호 존재, 제품명/모델명 일치, 인증상태 적합, 리콜 없음"

    if mismatch_fields and issues:
        message = "; ".join(mismatch_fields + issues)
        status = VALIDATION_MISMATCH if mismatch_fields else VALIDATION_INVALID

    return ValidationResult(
        row_number=record.row_number,
        validation_status=status,
        validation_message=message,
        cert_num_input=record.cert_num,
        product_name_input=record.product_name,
        model_name_input=record.model_name,
        cert_num_api=detail.cert_num,
        product_name_api=detail.product_name,
        model_name_api=detail.model_name,
        cert_status_api=detail.cert_status,
        recall_status_api=detail.recall_status,
        cert_date_api=detail.cert_date,
        cert_org_api=detail.cert_org,
        source_row=record.source_row,
    )


def write_report(path: Path, results: list[ValidationResult]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    base_fields = [
        "row_number",
        "validation_status",
        "validation_message",
        "cert_num_input",
        "product_name_input",
        "model_name_input",
        "cert_num_api",
        "product_name_api",
        "model_name_api",
        "cert_status_api",
        "recall_status_api",
        "cert_date_api",
        "cert_org_api",
    ]

    extra_fields: list[str] = []
    for result in results:
        for key in result.source_row:
            if key not in extra_fields and key != "details":
                extra_fields.append(key)

    fieldnames = base_fields + extra_fields

    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for result in results:
            row = {
                "row_number": result.row_number,
                "validation_status": result.validation_status,
                "validation_message": result.validation_message,
                "cert_num_input": result.cert_num_input,
                "product_name_input": result.product_name_input,
                "model_name_input": result.model_name_input,
                "cert_num_api": result.cert_num_api,
                "product_name_api": result.product_name_api,
                "model_name_api": result.model_name_api,
                "cert_status_api": result.cert_status_api,
                "recall_status_api": result.recall_status_api,
                "cert_date_api": result.cert_date_api,
                "cert_org_api": result.cert_org_api,
            }
            row.update({k: result.source_row.get(k, "") for k in extra_fields})
            writer.writerow(row)


def print_summary(results: list[ValidationResult]) -> None:
    counts: dict[str, int] = {}
    for result in results:
        counts[result.validation_status] = counts.get(result.validation_status, 0) + 1

    print("\n=== KC 인증 검증 결과 ===")
    for status in (
        VALIDATION_VALID,
        VALIDATION_MISMATCH,
        VALIDATION_INVALID,
        VALIDATION_NOT_FOUND,
        VALIDATION_SKIPPED,
    ):
        if status in counts:
            print(f"  {status}: {counts[status]}건")
    print(f"  TOTAL: {len(results)}건")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="RCC details 칼럼 기반 KC 인증정보 유효성 검증 (SafetyKorea searchPop API)",
    )
    parser.add_argument("--input", "-i", required=True, help="입력 CSV/Excel 경로")
    parser.add_argument("--output", "-o", required=True, help="검증 결과 CSV 리포트 경로")
    parser.add_argument("--details-column", default="details", help="인증정보 JSON/텍스트 칼럼명")
    parser.add_argument("--delay", type=float, default=0.3, help="API 요청 간 대기(초)")
    parser.add_argument(
        "--insecure",
        action="store_true",
        help="SSL 인증서 검증 건너뛰기 (macOS CERTIFICATE_VERIFY_FAILED 시)",
    )
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)

    if not input_path.exists():
        print(f"ERROR: 입력 파일이 없습니다: {input_path}", file=sys.stderr)
        return 1

    try:
        records = load_input_records(input_path, details_column=args.details_column)
    except RuntimeError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 1

    if not records:
        print("ERROR: 입력 데이터가 비어 있습니다.", file=sys.stderr)
        return 1

    results: list[ValidationResult] = []
    detail_cache: dict[str, KCCertDetail] = {}

    for index, record in enumerate(records):
        if not record.cert_num:
            results.append(validate_record(record, KCCertDetail("", "", "", "", "", "", "", {}, False)))
            continue

        if record.cert_num not in detail_cache:
            try:
                detail_cache[record.cert_num] = fetch_cert_detail(
                    record.cert_num,
                    insecure=args.insecure,
                )
            except RuntimeError as e:
                print(f"ERROR: {e}", file=sys.stderr)
                return 1
            if args.delay > 0 and index < len(records) - 1:
                time.sleep(args.delay)

        results.append(validate_record(record, detail_cache[record.cert_num]))

    write_report(output_path, results)
    print_summary(results)
    print(f"\n리포트 저장: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
